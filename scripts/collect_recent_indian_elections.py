from __future__ import annotations

import csv
import io
import json
import re
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import Request, urlopen

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "data" / "recent-indian-elections.json"
REPOSITORY = "thecont1/india-votes-data"
API_BASE = f"https://api.github.com/repos/{REPOSITORY}"
RAW_BASE = f"https://raw.githubusercontent.com/{REPOSITORY}"
USER_AGENT = "IndiaJourney public election data collector"
ECI_GE2024_BASE = "https://www.eci.gov.in/eci-backend/public/all_files/GE-2024-statistical-report"
ECI_LS2024_RESULTS_PATH = "GE-2024-statistical-report/33-Constituency-Wise-Detailed-Result.xls"
ECI_LS2024_RESULTS_URL = f"{ECI_GE2024_BASE}/33-Constituency-Wise-Detailed-Result.xls"
ECI_LS2024_PARTIES_URL = f"{ECI_GE2024_BASE}/3-List-Of-Political-Parties-Participated.xlsx"
# Surat was declared uncontested before polling, so ECI report 33 has no candidate rows for it.
SURAT_UNCONTESTED = ("Gujarat", "Surat", "MUKESHKUMAR CHANDRAKAANT DALAL", "BJP")
ASSEMBLY_FILE_PATTERN = re.compile(r"data/csv/20(?:22|23|24|25|26)Assembly-[A-Z]{2}\.csv")


def fetch_bytes(url: str, accept: str = "*/*") -> bytes:
    request = Request(url, headers={"User-Agent": USER_AGENT, "Accept": accept})
    with urlopen(request, timeout=60) as response:
        return response.read()


def source_snapshot() -> tuple[str, str, list[str]]:
    commit = json.loads(fetch_bytes(f"{API_BASE}/commits/main", "application/vnd.github+json"))
    commit_sha = commit["sha"]
    tree_sha = commit["commit"]["tree"]["sha"]
    tree = json.loads(fetch_bytes(
        f"{API_BASE}/git/trees/{tree_sha}?recursive=1",
        "application/vnd.github+json",
    ))
    files = sorted(
        path for path in (item.get("path", "") for item in tree.get("tree", []))
        if ASSEMBLY_FILE_PATTERN.fullmatch(path)
    )
    if not files:
        raise ValueError("The election data repository contains no supported recent Assembly CSV files")
    return commit_sha, commit["commit"]["committer"]["date"], files


def read_eci_party_names(content: bytes) -> dict[str, str]:
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(content), read_only=True).worksheets[0]
    names = {"IND": "Independent", "NOTA": "None of the Above"}
    for row in sheet.iter_rows(values_only=True):
        if len(row) > 4 and isinstance(row[0], (int, float)) and row[2] and row[4]:
            names.setdefault(str(row[2]).strip(), str(row[4]).strip().replace("\u2013", "-"))
    if len(names) < 100:
        raise ValueError("ECI party list workbook has too few parties")
    return names


def read_eci_results_rows(content: bytes) -> list[list]:
    import xlrd

    # ECI's published .xls has a benign OLE directory quirk that xlrd flags as corruption.
    sheet = xlrd.open_workbook(file_contents=content, ignore_workbook_corruption=True).sheet_by_index(0)
    header = [str(value).strip() for value in sheet.row_values(2)]
    if header[:3] != ["State Name", "PC Name", "Candidate Name"] or header[6] != "Party Name" or header[12] != "Total":
        raise ValueError(f"Unexpected ECI constituency-wise result layout: {header}")
    return [sheet.row_values(index) for index in range(3, sheet.nrows)]


def parse_lok_sabha_2024_rows(rows: list[list], party_names: dict[str, str]) -> list[dict]:
    seats: dict[tuple[str, str], list[dict]] = defaultdict(list)

    def add(state: str, constituency: str, candidate: str, party: str, votes: int | None) -> None:
        seats[(state.casefold(), constituency.casefold())].append({
            "year": 2024,
            "election_type": "Lok Sabha",
            "state_or_ut": state,
            "constituency_number": "",
            "constituency": constituency,
            "candidate": candidate,
            "party": party_names.get(party, party),
            "votes": votes,
            "result_status": "reported" if votes is not None else "uncontested",
            "electors": None,
            "source": ECI_LS2024_RESULTS_URL,
            "vote_share_method": None,
        })

    for row in rows:
        if len(row) < 13 or not str(row[1]).strip() or not str(row[2]).strip():
            continue
        try:
            votes = int(float(row[12]))
        except (TypeError, ValueError) as error:
            raise ValueError(f"Invalid ECI vote total for {row[1]}: {row[12]!r}") from error
        add(str(row[0]).strip(), str(row[1]).strip(), str(row[2]).strip(), str(row[6]).strip(), votes)
    if (SURAT_UNCONTESTED[0].casefold(), SURAT_UNCONTESTED[1].casefold()) not in seats:
        add(*SURAT_UNCONTESTED, None)
    return _with_vote_shares(seats.values())


def _with_vote_shares(seat_groups) -> list[dict]:
    observations = []
    for candidates in seat_groups:
        total_votes = sum(candidate["votes"] or 0 for candidate in candidates)
        for candidate in candidates:
            candidate["vote_share_percent"] = (
                round(candidate["votes"] / total_votes * 100, 4)
                if candidate["votes"] is not None and total_votes > 0
                else None
            )
            if candidate["votes"] is not None and total_votes > 0:
                candidate["vote_share_method"] = "calculated_from_listed_candidate_votes"
            observations.append(candidate)
    return observations


def _read_csv_rows(content: bytes, source_url: str) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig")
    if not text.strip():
        raise ValueError(f"Election CSV is empty: {source_url}")
    header = text.splitlines()[0]
    delimiter = ";" if header.count(";") > header.count(",") else ","
    rows = list(csv.DictReader(io.StringIO(text), delimiter=delimiter))
    if not rows:
        raise ValueError(f"Election CSV has no candidate rows: {source_url}")
    return [
        {str(key).strip().lstrip("\ufeff"): (value or "").strip() for key, value in row.items() if key}
        for row in rows
    ]


def _integer(value: str, field: str, source_url: str) -> int:
    try:
        result = int(value.replace(",", "").strip())
    except (ValueError, AttributeError) as error:
        raise ValueError(f"Invalid {field} value in {source_url}: {value!r}") from error
    if result < 0:
        raise ValueError(f"Negative {field} value in {source_url}: {value!r}")
    return result


def parse_recent_csv(
    content: bytes,
    relative_path: str,
    commit_sha: str,
    state_names: dict[str, str],
) -> list[dict]:
    source_url = f"{RAW_BASE}/{commit_sha}/{quote(relative_path, safe='/')}"
    rows = _read_csv_rows(content, source_url)
    required = {"election_year", "election_state", "constituency", "constituency_no",
                "candidate", "party", "evm_votes", "postal_votes"}
    missing = required - set(rows[0])
    if missing:
        raise ValueError(f"Election CSV is missing columns {sorted(missing)}: {source_url}")

    seats: dict[tuple[str, str, str, str], list[dict]] = defaultdict(list)
    for row in rows:
        candidate = row.get("candidate", "").strip()
        constituency = row.get("constituency", "").strip()
        if not candidate or not constituency or candidate.casefold() == "total":
            continue
        year = _integer(row["election_year"], "election year", source_url)
        if year not in range(2022, 2027):
            raise ValueError(f"Unexpected Assembly election year {year}: {source_url}")
        election_type = "State Assembly"
        state_code = row["election_state"].strip().upper()
        state = state_names.get(state_code, "")
        if not state:
            raise ValueError(f"Unknown ECI state code {state_code!r}: {source_url}")
        number = row["constituency_no"].strip()
        votes = (
            _integer(row["evm_votes"], "EVM votes", source_url)
            + _integer(row["postal_votes"], "postal votes", source_url)
        )
        seat_key = (str(year), election_type, state.casefold(), number or constituency.casefold())
        seats[seat_key].append({
            "year": year,
            "election_type": election_type,
            "state_or_ut": state,
            "constituency_number": number,
            "constituency": constituency,
            "candidate": candidate,
            "party": row.get("party", "").strip(),
            "votes": votes,
            "result_status": "reported",
            "electors": None,
            "source": source_url,
            "vote_share_method": None,
        })
    if not seats:
        raise ValueError(f"Election CSV contains no usable constituency results: {source_url}")
    return _with_vote_shares(seats.values())


def parse_state_names(content: bytes, commit_sha: str) -> dict[str, str]:
    source_url = f"{RAW_BASE}/{commit_sha}/data/states.csv"
    rows = _read_csv_rows(content, source_url)
    if not {"state_code", "state_name"} <= set(rows[0]):
        raise ValueError(f"State lookup is missing required columns: {source_url}")
    return {
        row["state_code"].strip().upper(): row["state_name"].strip()
        for row in rows
        if row.get("state_code") and row.get("state_name")
    }


def collect_recent_results() -> dict:
    commit_sha, commit_date, paths = source_snapshot()
    states_url = f"{RAW_BASE}/{commit_sha}/data/states.csv"
    state_names = parse_state_names(fetch_bytes(states_url), commit_sha)
    party_names = read_eci_party_names(fetch_bytes(ECI_LS2024_PARTIES_URL))
    observations: list[dict] = parse_lok_sabha_2024_rows(
        read_eci_results_rows(fetch_bytes(ECI_LS2024_RESULTS_URL)), party_names
    )
    source_files = [{"path": ECI_LS2024_RESULTS_PATH, "url": ECI_LS2024_RESULTS_URL}]
    source_index = {ECI_LS2024_RESULTS_URL: 0}
    for path in paths:
        source_url = f"{RAW_BASE}/{commit_sha}/{quote(path, safe='/')}"
        source_index[source_url] = len(source_files)
        source_files.append({"path": path, "url": source_url})
        observations.extend(parse_recent_csv(fetch_bytes(source_url), path, commit_sha, state_names))

    unique: dict[tuple, dict] = {}
    for row in observations:
        key = (
            row["year"],
            row["election_type"],
            row["state_or_ut"].casefold(),
            row["constituency"].casefold()
            if row["election_type"] == "Lok Sabha"
            else row["constituency_number"] or row["constituency"].casefold(),
            row["candidate"].casefold(),
            row["party"].casefold(),
            row["votes"],
        )
        unique[key] = row
    observations = list(unique.values())

    parliamentary_seats = {
        (row["state_or_ut"].casefold(), row["constituency"].casefold())
        for row in observations
        if row["year"] == 2024 and row["election_type"] == "Lok Sabha"
    }
    if len(parliamentary_seats) != 543:
        raise ValueError(
            f"Expected results for all 543 2024 Lok Sabha constituencies; found {len(parliamentary_seats)}"
        )
    for row in observations:
        row["source_file"] = source_index[row.pop("source")]
        row.pop("electors", None)
        if row.get("vote_share_method") == "calculated_from_listed_candidate_votes":
            row.pop("vote_share_method")
        if row.get("result_status") == "reported":
            row.pop("result_status")
    observations.sort(key=lambda row: (
        row["year"], row["election_type"], row["state_or_ut"], row["constituency_number"],
        -(row["votes"] or 0), row["candidate"].casefold(),
    ))
    return {
        "source": "2024 Lok Sabha: ECI GE 2024 statistical report 33 (Constituency Wise Detailed Result); 2022-2026 Assembly: ECI results normalized by the India Votes Data project. Candidate records link to the exact source file.",
        "source_repository": f"https://github.com/{REPOSITORY}",
        "source_commit": commit_sha,
        "source_commit_date": commit_date,
        "refresh_frequency": "Daily",
        "collection_status": "ok",
        "vote_share_method": "Calculated from the listed candidate vote totals (EVM + postal) for that constituency; uncontested candidates without a reported vote count have no calculated share.",
        "source_files": source_files,
        "observation_count": len(observations),
        "lok_sabha_2024_constituency_count": len(parliamentary_seats),
        "assembly_election_years": sorted({
            row["year"] for row in observations if row["election_type"] == "State Assembly"
        }),
        "observations": observations,
        "generated_at_utc": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
    }


def main() -> int:
    try:
        payload = collect_recent_results()
    except (HTTPError, URLError, TimeoutError, KeyError, ValueError, json.JSONDecodeError) as error:
        print(f"Unable to refresh recent Indian election results: {error}", file=sys.stderr)
        return 1
    try:
        previous = json.loads(OUTPUT.read_text(encoding="utf-8"))
    except FileNotFoundError:
        previous = {}
    except json.JSONDecodeError as error:
        print(f"Unable to read existing recent election snapshot at {OUTPUT}: {error}", file=sys.stderr)
        return 1
    if previous.get("source_commit") == payload["source_commit"]:
        payload["generated_at_utc"] = previous.get("generated_at_utc", payload["generated_at_utc"])
    serialized = json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(serialized, encoding="utf-8")
    (ROOT / "docs" / "data" / OUTPUT.name).write_text(serialized, encoding="utf-8")
    print(json.dumps({
        "collection_status": payload["collection_status"],
        "source_files": len(payload["source_files"]),
        "candidate_observations": payload["observation_count"],
        "2024_lok_sabha_constituencies": payload["lok_sabha_2024_constituency_count"],
        "assembly_years": payload["assembly_election_years"],
        "source_commit": payload["source_commit"],
    }))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
